"""Moteur de calcul des plans d'amortissement (vente à tempérament) – réplique de l'onglet Moteur du classeur Excel."""
import math
import numpy as np
import pandas as pd
import openpyxl
from dateutil.relativedelta import relativedelta


def xround(x, nd=0):
    """ROUND d'Excel (arrondi commercial, .5 vers le haut en valeur absolue)."""
    m = 10 ** nd
    return math.copysign(math.floor(abs(x) * m + 0.5) / m, x)


def pmt(rate, n, pv):
    """Mensualité constante (équivalent PMT(rate, n, -pv))."""
    return pv / n if rate == 0 else pv * rate / (1 - (1 + rate) ** (-n))


def _tbl(ws, header_row, first_row, last_row, ncols):
    heads = [ws.cell(header_row, c).value for c in range(1, ncols + 1)]
    rows = [[ws.cell(r, c).value for c in range(1, ncols + 1)] for r in range(first_row, last_row + 1)]
    return pd.DataFrame(rows, columns=heads)


def lire_classeur(chemin):
    """Lit les onglets de SAISIE du classeur Excel (les formules sont ignorées).
    `chemin` peut être un chemin de fichier ou un fichier en mémoire (upload Streamlit)."""
    wb = openpyxl.load_workbook(chemin, data_only=True)
    p = wb["Parametres"]
    params = dict(
        tiao=float(p["B4"].value), frais=float(p["B5"].value),
        majoration=float(p["B7"].value), base_jours=int(p["B9"].value),
        date_ref=pd.Timestamp(p["B10"].value), devise=p["B11"].value or "FCFA",
        prix_parking=float(p["B13"].value or 0),
    )
    params["taux_base"] = params["tiao"] + params["frais"]
    params["taux_pen"] = params["taux_base"] + params["majoration"]
    params["groupes"] = [p.cell(r, 5).value for r in range(4, 14) if p.cell(r, 5).value]

    prod = _tbl(wb["Produits"], 5, 6, 405, 13).iloc[:, :13]
    prod.columns = ["adressage", "categorie", "niveau", "type", "num_appt", "lot_edd",
                    "superficie", "prix_m2", "prix_total", "alloc_parking", "loc_parking", "emplacement", "observation"]
    prod = prod[prod.adressage.notna()].copy()

    def prix_retenu(r):
        if pd.notna(r.prix_total): return float(r.prix_total)
        if pd.notna(r.prix_m2) and pd.notna(r.superficie): return xround(r.superficie * r.prix_m2)
        return params["prix_parking"] if r.categorie == "Parking" else 0.0
    prod["prix_ht"] = prod.apply(prix_retenu, axis=1)

    acq = _tbl(wb["Acquereurs"], 5, 6, 45, 9).iloc[:, :9]
    acq.columns = ["id", "nom", "groupe", "bien", "prix_bien", "versement_initial", "duree", "date_1ere", "taux_specifique"]
    acq = acq[acq.id.notna()].drop(columns=["bien", "prix_bien"]).copy()
    acq["date_1ere"] = pd.to_datetime(acq["date_1ere"])
    acq["versement_initial"] = acq["versement_initial"].fillna(0).astype(float)
    acq["duree"] = acq["duree"].astype(int)

    ven = _tbl(wb["Ventes"], 4, 5, 404, 3)
    ven.columns = ["id", "adressage", "ajustement"]
    ven = ven[ven.adressage.notna()].copy()
    ven["ajustement"] = ven["ajustement"].fillna(0).astype(float)

    pai = _tbl(wb["Paiements"], 4, 5, 1004, 4)
    pai.columns = ["id", "echeance", "date", "montant"]
    pai = pai[pai.id.notna()].copy(); pai["date"] = pd.to_datetime(pai["date"])

    rem = _tbl(wb["Remb_Anticipes"], 4, 5, 504, 3)
    rem.columns = ["id", "date", "montant"]
    rem = rem[rem.id.notna()].copy(); rem["date"] = pd.to_datetime(rem["date"])
    return dict(params=params, produits=prod, acquereurs=acq, ventes=ven, paiements=pai, remb=rem)


def controles(d):
    """Contrôles d'intégrité (doublons, ID inconnus…)."""
    v = d["ventes"].merge(d["produits"][["adressage", "categorie"]], on="adressage", how="left")
    msgs = []
    for _, r in v.iterrows():
        if pd.isna(r.categorie): msgs.append(f"Produit inconnu : {r.adressage}")
        elif r.id not in set(d["acquereurs"].id): msgs.append(f"ID acquéreur inconnu en Ventes : {r.id}")
    for a in v.adressage[v.adressage.duplicated()].unique(): msgs.append(f"Produit vendu en doublon : {a}")
    for i in d["paiements"].id[~d["paiements"].id.isin(d["acquereurs"].id)].unique(): msgs.append(f"Paiement : ID inconnu {i}")
    for i in d["remb"].id[~d["remb"].id.isin(d["acquereurs"].id)].unique(): msgs.append(f"Remb. anticipé : ID inconnu {i}")
    return msgs


def calculer(d):
    """Retourne (synthèse par acquéreur, échéancier complet, ventes détaillées)."""
    P = d["params"]; ref = P["date_ref"]
    ven = d["ventes"].merge(d["produits"][["adressage", "categorie", "superficie", "prix_ht"]], on="adressage", how="left")
    ven["prix_net"] = ven.prix_ht.fillna(0) + ven.ajustement
    ven["is_parking"] = ven.categorie.eq("Parking")

    lignes, synth = [], []
    for a in d["acquereurs"].itertuples():
        vv = ven[ven.id == a.id]
        prix = vv.prix_net.sum()
        capital = max(0.0, prix - a.versement_initial)
        taux = a.taux_specifique if pd.notna(a.taux_specifique) else P["taux_base"]
        r = taux / 12
        mens = 0.0 if capital == 0 else float(xround(pmt(r, a.duree, capital)))
        pai = d["paiements"][d["paiements"].id == a.id].groupby("echeance").montant.sum()
        rem = d["remb"][d["remb"].id == a.id]
        cap_fin = None
        for n in range(1, a.duree + 1):
            ech = a.date_1ere + relativedelta(months=n - 1)
            cap_deb = capital if n == 1 else cap_fin
            interets = float(xround(cap_deb * r))
            if cap_deb < 0.5: due = 0.0
            elif n == a.duree: due = cap_deb + interets
            else: due = min(mens, cap_deb + interets)
            amorti = due - interets
            debut = pd.Timestamp(0) if n == 1 else ech - relativedelta(months=1)
            rb = rem[(rem.date > debut) & (rem.date <= ech)].montant.sum()
            anticipe = max(0.0, min(cap_deb - amorti, rb))
            cap_fin = cap_deb - amorti - anticipe
            paye = float(pai.get(n, 0.0))
            impaye = max(0.0, due - paye) if (due > 0 and ech <= ref) else 0.0
            jours = (ref - ech).days if impaye > 0 else 0
            penal = float(xround(impaye * P["taux_pen"] * jours / P["base_jours"]))
            if due == 0: st = "Soldé"
            elif paye >= due: st = "Payé"
            elif ech > ref: st = "À venir"
            else: st = "Partiel" if paye > 0 else "Impayé"
            lignes.append(dict(id=a.id, nom=a.nom, groupe=a.groupe, n=n, date_echeance=ech, capital_debut=cap_deb,
                               interets=interets, mensualite_due=due, capital_amorti=amorti, remb_anticipe=anticipe,
                               capital_fin=cap_fin, paye=paye, impaye=impaye, jours_retard=jours, penalite=penal, statut=st))
        e = pd.DataFrame([l for l in lignes if l["id"] == a.id])
        echu = e.loc[e.date_echeance <= ref, "mensualite_due"].sum()
        crd = max(0.0, capital - e.loc[e.date_echeance <= ref, ["capital_amorti", "remb_anticipe"]].sum().sum())
        imp, pen = e.impaye.sum(), e.penalite.sum()
        statut = "Soldé" if crd <= 0.5 else ("En retard" if imp > 0.5 else "À jour")
        nb = int((e.mensualite_due > 0).sum())
        synth.append(dict(id=a.id, nom=a.nom, groupe=a.groupe, nb_lots=int((~vv.is_parking).sum()), nb_parkings=int(vv.is_parking.sum()),
                          surface_m2=vv.superficie.fillna(0).sum(), prix_bien=prix, versement_initial=a.versement_initial,
                          capital_finance=capital, taux_annuel=taux, mensualite=mens,
                          interets_totaux=e.interets.sum(), nb_echeances=nb,
                          date_fin=a.date_1ere + relativedelta(months=nb - 1) if nb else pd.NaT,
                          echu_a_date=echu, total_paye=d["paiements"].loc[d["paiements"].id == a.id, "montant"].sum(),
                          remb_anticipes=rem[rem.date <= ref].montant.sum(), capital_restant_du=crd,
                          impayes=imp, penalites=pen, total_du=imp + pen,
                          taux_recouvrement=(1 - imp / echu) if echu else np.nan, statut=statut))
    return pd.DataFrame(synth), pd.DataFrame(lignes), ven
